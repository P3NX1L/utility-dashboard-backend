"""Validated full-snapshot imports. Preview never changes the active ledger."""
import base64
from collections import Counter, defaultdict
from contextlib import redirect_stdout
from datetime import date
from decimal import Decimal
import hashlib
import io
import json
from pathlib import Path
import re
import tempfile
from uuid import UUID, uuid4, uuid5
import zipfile

from fastapi import HTTPException
from pydantic import BaseModel, ConfigDict, Field

from dashboard_logic import cents
from workflow import require_role, rpc

NAMESPACE = UUID('c5cfec7b-726a-46cd-a8b9-87adf39e072f')
MAX_BYTES = 5 * 1024 * 1024
MAX_ROWS = 20000
PUBLIC_IMPORT_COLUMNS = ('id,filename,status,created_at,activated_at,source_exported_at,validation_report,'
                         'revision:scope->revision,base_import_id:scope->base_import_id,'
                         'pending_matches:scope->pending_matches')


class UploadRequest(BaseModel):
    model_config = ConfigDict(extra='forbid')
    filename: str = Field(min_length=1, max_length=180)
    content_base64: str = Field(max_length=7 * 1024 * 1024)


class ActivateRequest(BaseModel):
    model_config = ConfigDict(extra='forbid')
    expected_revision: int = Field(ge=0)
    acknowledge_warnings: bool = False


class MatchRequest(BaseModel):
    model_config = ConfigDict(extra='forbid')
    expected_revision: int = Field(ge=0)
    source_row_id: UUID
    bill_id: UUID | None = None  # null explicitly means a new transaction


def uid(*parts):
    return str(uuid5(NAMESPACE, json.dumps(parts)))


def normalize(value):
    return re.sub(r'\s+', ' ', str(value or '')).strip().casefold()


def usage(description):
    matches = re.findall(r'\busage\s*[:=]?\s*(\d+(?:\.\d+)?)\b', description or '', re.I)
    # Preserve the number, never invent gallons or another physical unit.
    return matches[0] if len(matches) == 1 else None


def validate_workbook(content):
    if not content or len(content) > MAX_BYTES:
        raise ValueError('Choose a nonempty .xlsx file no larger than 5 MB.')
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        entries = archive.infolist()
        if len(entries) > 1000 or sum(e.file_size for e in entries) > 50 * 1024 * 1024:
            raise ValueError('Workbook expands beyond the supported size.')
        if any('vbaproject' in e.filename.lower() or 'externallinks/' in e.filename.lower() for e in entries):
            raise ValueError('Macros and external workbook links are not accepted.')
    # Read formulas separately; pandas can otherwise silently read stale cached values.
    from openpyxl import load_workbook
    try:
        book = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
    except Exception:
        raise ValueError('Workbook cannot be read. Upload a valid .xlsx vendor-ledger export.') from None
    try:
        if sum(s.max_row or 0 for s in book.worksheets) > MAX_ROWS:
            raise ValueError('Workbook exceeds 20,000 source rows.')
        for sheet in book:
            if (sheet.max_column or 0) > 100:
                raise ValueError('Workbook has too many columns.')
            for row in sheet.iter_rows():
                if any(cell.data_type == 'f' for cell in row):
                    raise ValueError('Export formulas as values before uploading; cached formula results cannot be trusted.')
    finally:
        book.close()


def parse_workbook(content, filename):
    validate_workbook(content)
    from prepare_ledger import run
    with tempfile.TemporaryDirectory(prefix='utility-import-') as tmp:
        source = Path(tmp) / 'source.xlsx'
        source.write_bytes(content)
        with redirect_stdout(io.StringIO()):
            run(source, Path(tmp) / 'parsed')
        def read(name):
            return json.loads((Path(tmp) / 'parsed' / f'{name}.json').read_text())
        records, raw, report = read('transactions'), read('raw_rows'), read('quality_report')
    if not records or report['rejected_rows'] or any(not r['account_number'] for r in records):
        raise ValueError('No import created: transactions are missing, rejected rows exist, or an account could not be parsed.')
    if len(records) > MAX_ROWS:
        raise ValueError('Workbook has too many transactions.')
    for r in records:
        r['source_file'] = filename
    return records, raw, report


def identity(row):
    return (row['account_id'], row['category'], row['due_date'], normalize(row['description']),
            row['transaction_type'], row['reading_type'])


def financial_signature(row):
    return identity(row) + (cents(row['paid_amount']), cents(row['unpaid_amount']),
                           row.get('paid_date'), normalize(row.get('payment_reference')),
                           row.get('service_start'), row.get('service_end'), row.get('period_status'))


def build_preview(records, raw, report, active, old_rows, properties=None, accounts=None):
    """Match exact unique rows, then unique equal-total payment/status changes.
    Ambiguous matches are never silently attached to an existing investigation.
    """
    from load_initial_import import build_payload
    import_id = str(uuid4())
    payload = build_payload(records, raw, report, import_id)
    # Reuse canonical property/account IDs even if future naming rules change.
    prop_map = {p['property_key']: p['id'] for p in properties or []}
    account_map = {(a['property_id'], a['provider'], a['account_number']): a['id'] for a in accounts or []}
    prop_ids = {}
    for p in payload['properties']:
        original = p['id']; p['id'] = prop_map.get(p['property_key'], original); prop_ids[original] = p['id']
    acct_ids = {}
    for a in payload['utility_accounts']:
        original = a['id']; a['property_id'] = prop_ids[a['property_id']]
        a['id'] = account_map.get((a['property_id'], a['provider'], a['account_number']), original)
        acct_ids[original] = a['id']
    for b in payload['bills']:
        b['account_id'] = acct_ids[b['account_id']]
    bill_accounts = {b['id']: b['account_id'] for b in payload['bills']}
    source_records = {(r['source_sheet'], r['source_row']): r for r in records}
    sources = {r['id']: r for r in payload['import_rows']}
    old_identity = defaultdict(list)
    old_day = defaultdict(list)
    for row in old_rows:
        old_identity[identity(row)].append(row)
        old_day[(row['account_id'], row['category'], row['due_date'])].append(row)
    new_identity = Counter()
    new_financial = Counter()
    for v in payload['bill_versions']:
        v['account_id'] = bill_accounts[v['bill_id']]
        new_identity[identity(v)] += 1
        new_financial[financial_signature(v)] += 1
    matched, pending, counts = set(), [], Counter()
    for v in payload['bill_versions']:
        source = sources[v['source_row_id']]
        record = source_records[(source['sheet_name'], source['row_number'])]
        v['usage_value'] = usage(v['description'])
        v['usage_unit'] = None
        candidates = old_identity[identity(v)]
        same_identity = bool(candidates)
        if not candidates:
            candidates = old_day[(v['account_id'], v['category'], v['due_date'])]
        exact = [o for o in candidates if financial_signature(o) == financial_signature(v)]
        chosen = None
        if len(exact) == 1 and new_financial[financial_signature(v)] == 1:
            chosen = exact[0]; kind = 'unchanged'
        elif same_identity and len(candidates) == 1 and new_identity[identity(v)] == 1 and (
                cents(candidates[0]['paid_amount']) + cents(candidates[0]['unpaid_amount']) ==
                cents(v['paid_amount']) + cents(v['unpaid_amount'])):
            chosen = candidates[0]; kind = 'updated'
        elif not candidates:
            kind = 'new'
        else:
            kind = 'uncertain'
            pending.append({'source_row_id': v['source_row_id'], 'sheet': source['sheet_name'],
                            'row': source['row_number'], 'description': v['description'],
                            'total_cents': record['total_cents'], 'property_name': record['property_label'],
                            'candidates': [{'bill_id': o['bill_id'], 'bill_version_id': o['id'],
                                            'source': {'sheet': o['source_sheet'], 'row': o['source_row']},
                                            'paid_cents': cents(o['paid_amount']), 'unpaid_cents': cents(o['unpaid_amount']),
                                            'description': o['description']} for o in candidates]})
        if chosen:
            if chosen['bill_id'] in matched:
                raise ValueError('One existing bill cannot match two imported rows.')
            matched.add(chosen['bill_id']); v['bill_id'] = chosen['bill_id']
            v['match_method'] = 'exact' if kind == 'unchanged' else 'rule_based'
        v['id'] = uid('version', import_id, v['source_row_id'])
        v['match_evidence'].update(change_kind=kind, account_id=v['account_id'],
                                   basis='unique_identity_equal_total' if kind == 'updated' else kind)
        v.pop('account_id')
        counts[kind] += 1
    # Full snapshots replace the active view, not historical records.
    report = dict(report)
    report['matching'] = dict(counts)
    report['matching'].update(removed_from_snapshot=len({o['bill_id'] for o in old_rows} - matched),
                               previous_transaction_count=len(old_rows))
    warnings = []
    if any(r['difference_paid_cents'] or r['difference_unpaid_cents'] for r in report['reconciliation']):
        warnings.append('Source control totals differ from transaction rows. No balancing transactions will be added.')
    if report['duplicate_groups']:
        warnings.append('Potential duplicate groups are retained in full and require review.')
    if report['matching']['removed_from_snapshot']:
        warnings.append('Some old transactions are absent or not yet matched. Activation replaces the full dashboard snapshot; history remains stored.')
    report['warnings'] = warnings
    exported = None
    return {'id': import_id, 'report': report, 'payload': payload, 'pending_matches': pending,
            'base_import_id': active['id'], 'base_bill_ids': sorted({o['bill_id'] for o in old_rows}),
            'revision': 0, 'source_exported_at': exported}


def public_preview(item):
    scope = item.get('scope') or {}
    if not scope:
        scope = {key: item.get(key) for key in ('revision', 'base_import_id', 'pending_matches')}
    report = item.get('validation_report') or {}
    return {'import_id': item['id'], 'filename': item['filename'], 'status': item['status'],
            'created_at': item.get('created_at'), 'activated_at': item.get('activated_at'),
            'source_exported_at': item.get('source_exported_at'),
            'revision': scope.get('revision') or 0, 'base_import_id': scope.get('base_import_id'),
            'validation_report': report, 'pending_matches': scope.get('pending_matches') or [],
            'transaction_count': report.get('transaction_count'), 'property_count': report.get('property_count'),
            'total_cents': report.get('totals_cents', {}).get('total_cents')}


def stage_upload(db, member, request, dataset):
    require_role(member, ('admin', 'reviewer'))
    filename = Path(request.filename.replace('\\', '/')).name
    if not filename.lower().endswith('.xlsx') or any(ord(c) < 32 for c in filename):
        raise HTTPException(422, 'Only .xlsx exports are accepted.')
    try:
        content = base64.b64decode(request.content_base64, validate=True)
        validate_workbook(content)
    except (ValueError, zipfile.BadZipFile, KeyError, OSError) as exc:
        raise HTTPException(422, str(exc)) from None
    file_hash = hashlib.sha256(content).hexdigest()
    existing = db.table('imports').select(PUBLIC_IMPORT_COLUMNS).eq('file_sha256', file_hash).limit(1).execute().data
    if existing:
        return {**public_preview(existing[0]), 'already_uploaded': True}
    try:
        records, raw, report = parse_workbook(content, filename)
        from workflow import read_all
        properties = read_all(db.table('properties').select('id,property_key,name'))
        accounts = read_all(db.table('utility_accounts').select('id,property_id,provider,account_number'))
        preview = build_preview(records, raw, report, *dataset, properties, accounts)
    except (ValueError, TypeError, KeyError) as exc:
        raise HTTPException(422, str(exc)) from None
    bucket = dataset[0].get('scope', {}).get('bucket') or 'utility-source-files'
    storage_path = f"imports/{preview['id']}/source.xlsx"
    db.storage.from_(bucket).upload(storage_path, content, {'content-type': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', 'upsert': 'false'})
    item = {'id': preview['id'], 'filename': filename, 'storage_path': storage_path,
            'file_sha256': file_hash, 'source_row_count': len(raw),
            'source_exported_at': preview['source_exported_at'], 'validation_report': preview['report'],
            'scope': {'mode': 'full_snapshot', 'bucket': bucket, 'base_import_id': preview['base_import_id'],
                      'revision': 0, 'base_bill_ids': preview['base_bill_ids'],
                      'payload': preview['payload'], 'pending_matches': preview['pending_matches']}}
    stored = rpc(db, 'stage_utility_import', {'p_actor_id': member['user_id'], 'p_import': item})
    return public_preview(stored)


def payload_rows(payload):
    properties = {r['id']: r for r in payload['properties']}
    accounts = {r['id']: r for r in payload['utility_accounts']}
    bills = {r['id']: r for r in payload['bills']}
    sources = {r['id']: r for r in payload['import_rows']}
    # Matched bills no longer have their draft ID. Use explicit account evidence.
    result = []
    for v in payload['bill_versions']:
        a = accounts[v['match_evidence']['account_id']]
        p = properties[a['property_id']]
        s = sources[v['source_row_id']]
        result.append({**v, 'account_id': a['id'], 'account_number': a['account_number'],
                       'property_id': p['id'], 'property_name': p['name'],
                       'source_sheet': s['sheet_name'], 'source_row': s['row_number'],
                       'source_filename': payload.get('filename'),
                       'duration_days': (date.fromisoformat(v['service_end']) - date.fromisoformat(v['service_start'])).days
                       if v['service_start'] and v['service_end'] else None})
    return result


def activate_import(db, member, import_id, request):
    require_role(member, ('admin',))
    found = db.table('imports').select('*').eq('id', str(import_id)).limit(1).execute().data
    if not found:
        raise HTTPException(404, 'Import not found.')
    item = found[0]
    if item['status'] == 'active':
        return {'import_id': item['id'], 'already_active': True}
    scope = item['scope']
    if scope.get('revision') != request.expected_revision or scope.get('pending_matches'):
        raise HTTPException(409, 'Refresh the preview and resolve every uncertain match before activation.')
    if item['validation_report'].get('warnings') and not request.acknowledge_warnings:
        raise HTTPException(422, 'Acknowledge the source-total, duplicate and full-snapshot warnings before activation.')
    from deviation_logic import evaluate, METHOD_VERSION, PARAMETERS
    assessments = evaluate(payload_rows(scope['payload']))
    return rpc(db, 'activate_utility_import', {
        'p_import_id': str(import_id), 'p_actor_id': member['user_id'],
        'p_expected_revision': request.expected_revision, 'p_acknowledge_warnings': request.acknowledge_warnings,
        'p_method': METHOD_VERSION, 'p_parameters': PARAMETERS,
        'p_assessments': sorted(assessments, key=lambda r: r['bill_version_id'])})
